import json
import random
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from functools import wraps
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.conf import settings
from django.core.paginator import Paginator
from django.contrib import messages
from django.http import HttpResponse
from pypdf import PdfReader
import docx
from pptx import Presentation
from openai import OpenAI
from django.db.models import Avg
from django.contrib.auth.models import User, Group

from .forms import PreguntaForm, SubirArchivoForm, UsuarioForm, EditarUsuarioForm, CambiarPasswordForm
from .models import Pregunta, Opcion, Especialidad, Curso, Anio, Seccion, ResultadoExamen, ExamenAsignado

def docente_required(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if request.user.is_superuser or request.user.groups.filter(name='Docentes').exists():
            return view_func(request, *args, **kwargs)
        messages.error(request, "Acceso denegado. Exclusivo para Docentes.")
        return redirect('home')
    return _wrapped_view

@login_required
def home(request):
    es_docente = request.user.is_superuser or request.user.groups.filter(name='Docentes').exists()
    total_preguntas = Pregunta.objects.count()
    total_examenes_asignados = ExamenAsignado.objects.count()
    total_resultados = ResultadoExamen.objects.count()
    promedio_general = ResultadoExamen.objects.aggregate(Avg('nota'))['nota__avg'] or 0
    
    ultimas_notas = ResultadoExamen.objects.all().order_by('-fecha_rendido')[:5] if es_docente else ResultadoExamen.objects.filter(alumno=request.user).order_by('-fecha_rendido')[:5]

    tipos_data = {
        'labels': ['Opción Múltiple', 'Correlaciones', 'V/F', 'Completar'],
        'values': [Pregunta.objects.filter(tipo='OM').count(), Pregunta.objects.filter(tipo__startswith='CORR').count(), Pregunta.objects.filter(tipo='VF').count(), Pregunta.objects.filter(tipo='COMP').count()]
    }

    examenes_pendientes = []
    if not es_docente:
        rendidos_ids = ResultadoExamen.objects.filter(alumno=request.user, examen_asignado__isnull=False).values_list('examen_asignado_id', flat=True)
        examenes_pendientes = ExamenAsignado.objects.filter(activo=True).exclude(id__in=rendidos_ids)

    return render(request, 'core/home.html', {
        'es_docente': es_docente, 'especialidades': Especialidad.objects.all(), 'cursos': Curso.objects.all(), 'anios': Anio.objects.all(),
        'examenes_pendientes': examenes_pendientes, 'total_preguntas': total_preguntas, 'total_examenes_asignados': total_examenes_asignados,
        'total_resultados': total_resultados, 'promedio_general': round(promedio_general, 1), 'ultimas_notas': ultimas_notas, 'tipos_data': tipos_data
    })

@login_required
@docente_required
def gestionar_usuarios(request):
    if not request.user.is_superuser:
        messages.error(request, "Solo el Administrador puede gestionar usuarios.")
        return redirect('home')
    usuarios = User.objects.all().order_by('username')
    return render(request, 'core/gestion_usuarios.html', {'usuarios': usuarios})

@login_required
@docente_required
def crear_usuario(request):
    if not request.user.is_superuser: return redirect('home')
    if request.method == 'POST':
        form = UsuarioForm(request.POST)
        if form.is_valid():
            usuario = form.save(commit=False)
            usuario.save()
            rol = form.cleaned_data['rol']
            grupo, _ = Group.objects.get_or_create(name=rol)
            usuario.groups.add(grupo)
            messages.success(request, f"Usuario {usuario.username} creado como {rol}.")
            return redirect('gestionar_usuarios')
    else: form = UsuarioForm()
    return render(request, 'core/crear_usuario.html', {'form': form})

@login_required
@docente_required
def editar_usuario(request, usuario_id):
    if not request.user.is_superuser: return redirect('home')
    usuario = get_object_or_404(User, id=usuario_id)
    if request.method == 'POST':
        form = EditarUsuarioForm(request.POST, instance=usuario)
        if form.is_valid():
            user_edit = form.save(commit=False)
            user_edit.is_active = form.cleaned_data['activo']
            nuevo_rol = form.cleaned_data['rol']
            if nuevo_rol == 'Administrador':
                user_edit.is_superuser = True
                user_edit.is_staff = True
            else:
                user_edit.is_superuser = False
                user_edit.is_staff = False
                user_edit.groups.clear()
                grupo, _ = Group.objects.get_or_create(name=nuevo_rol)
                user_edit.groups.add(grupo)
            user_edit.save()
            messages.success(request, f"Usuario {usuario.username} actualizado correctamente.")
            return redirect('gestionar_usuarios')
    else:
        form = EditarUsuarioForm(instance=usuario)
    return render(request, 'core/editar_usuario.html', {'form': form, 'usuario': usuario})

@login_required
@docente_required
def eliminar_usuario(request, usuario_id):
    if not request.user.is_superuser: return redirect('home')
    usuario = get_object_or_404(User, id=usuario_id)
    if usuario == request.user:
        messages.error(request, "No puedes eliminar tu propio usuario administrador.")
    else:
        usuario.delete()
        messages.success(request, "Usuario eliminado correctamente.")
    return redirect('gestionar_usuarios')

@login_required
@docente_required
def cargar_usuarios_excel(request):
    if not request.user.is_superuser: return redirect('home')
    if request.method == 'POST':
        archivo = request.FILES.get('archivo_excel')
        if not archivo:
            messages.error(request, "Por favor, selecciona un archivo Excel antes de presionar procesar.")
            return render(request, 'core/cargar_usuarios.html')
        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
            hoja = wb.active
            grupo_alumnos, _ = Group.objects.get_or_create(name='Alumnos')
            creados, errores = 0, 0
            for row in hoja.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]: continue
                username = str(row[0]).strip() if row[0] else ''
                first_name = str(row[1]).strip() if len(row) > 1 and row[1] else ''
                last_name = str(row[2]).strip() if len(row) > 2 and row[2] else ''
                email = str(row[3]).strip() if len(row) > 3 and row[3] else ''
                if not username: continue
                if User.objects.filter(username=username).exists():
                    errores += 1
                    continue
                user = User.objects.create_user(username=username, email=email, password='Citen2024*')
                user.first_name = first_name
                user.last_name = last_name
                user.save()
                user.groups.add(grupo_alumnos)
                creados += 1
            if creados > 0: messages.success(request, f"¡Éxito! Se crearon {creados} alumnos. Contraseña por defecto: Citen2024*")
            if errores > 0: messages.warning(request, f"Se omitieron {errores} usuarios porque ya existían.")
            return redirect('gestionar_usuarios')
        except Exception as e:
            messages.error(request, f"Error al leer el archivo Excel: {str(e)}")
            return render(request, 'core/cargar_usuarios.html')
    return render(request, 'core/cargar_usuarios.html')

@login_required
def cambiar_password(request):
    if request.method == 'POST':
        form = CambiarPasswordForm(user=request.user, data=request.POST)
        if form.is_valid():
            user = form.save()
            if hasattr(user, 'profile'):
                user.profile.must_change_password = False
                user.profile.save()
            messages.success(request, "Tu contraseña ha sido actualizada correctamente.")
            return redirect('home')
    else:
        form = CambiarPasswordForm(user=request.user)
    return render(request, 'core/cambiar_password.html', {'form': form})

@login_required
@docente_required
def registrar_pregunta(request):
    if request.method == 'POST':
        form = PreguntaForm(request.POST)
        if form.is_valid():
            pregunta = form.save(commit=False)
            pregunta.creado_por, pregunta.nivel = request.user, request.POST.get('nivel', 'FACIL')
            tipo = form.cleaned_data['tipo']
            if tipo in ['CORR', 'CORR_CAUSA', 'CORR_DEF']:
                pares = {request.POST.get(f'corr_a_{i}'): request.POST.get(f'corr_b_{i}') for i in range(1, 6) if request.POST.get(f'corr_a_{i}') and request.POST.get(f'corr_b_{i}')}
                pregunta.contenido = {"pares": pares}
            elif tipo == 'VF': pregunta.contenido = {"correcta": request.POST.get('vf_respuesta') == 'true'}
            elif tipo == 'COMP': pregunta.contenido = {"palabra": request.POST.get('comp_palabra', '')}
            elif tipo == 'LIBRE': pregunta.contenido = {"respuesta_modelo": request.POST.get('libre_respuesta', '')}
            pregunta.save()
            if tipo == 'OM':
                opciones = [request.POST.get('om_a', ''), request.POST.get('om_b', ''), request.POST.get('om_c', ''), request.POST.get('om_d', '')]
                for idx, texto in enumerate(opciones):
                    if texto: Opcion.objects.create(pregunta=pregunta, texto=texto, es_correcta=(idx == int(request.POST.get('om_correcta', 0))), orden=idx)
            messages.success(request, "Pregunta registrada exitosamente.")
            return redirect('home')
    else: form = PreguntaForm()
    return render(request, 'core/registrar_pregunta.html', {'form': form})

def extraer_texto(archivo):
    texto = ""
    nombre = archivo.name.lower()
    try:
        if nombre.endswith('.pdf'):
            reader = PdfReader(archivo)
            for page in reader.pages: texto += (page.extract_text() or "") + "\n"
        elif nombre.endswith('.docx'):
            doc = docx.Document(archivo)
            for para in doc.paragraphs: texto += para.text + "\n"
        elif nombre.endswith('.pptx'):
            prs = Presentation(archivo)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"): texto += shape.text + "\n"
        else: raise ValueError("Formato no soportado.")
    except Exception as e: raise ValueError(f"Error al leer el archivo: {str(e)}")
    return texto

@login_required
@docente_required
def subir_material(request):
    if request.method == 'POST':
        form = SubirArchivoForm(request.POST, request.FILES)
        if form.is_valid():
            archivo = request.FILES['archivo']
            try:
                texto = extraer_texto(archivo)
                if not texto.strip(): return render(request, 'core/subir_material.html', {'form': form, 'error': "El archivo no tiene texto legible."})
                
                nuevo_curso_nombre = form.cleaned_data.get('nuevo_curso')
                curso_seleccionado = form.cleaned_data.get('curso')
                especialidad = form.cleaned_data.get('especialidad')
                
                if nuevo_curso_nombre and nuevo_curso_nombre.strip() != "":
                    curso_final, _ = Curso.objects.get_or_create(nombre=nuevo_curso_nombre.strip(), especialidad=especialidad)
                elif curso_seleccionado:
                    curso_final = curso_seleccionado
                else:
                    return render(request, 'core/subir_material.html', {'form': form, 'error': "Debes seleccionar un curso existente o escribir uno nuevo."})

                anio_seleccionado = form.cleaned_data['anio']
                seccion_seleccionada = form.cleaned_data['seccion']
                nombre_examen = form.cleaned_data.get('nombre_examen') or f"Examen de {curso_final.nombre}"

                texto_recortado = texto[:6000] 
                api_key = os.getenv('GROQ_API_KEY')
                if not api_key: return render(request, 'core/subir_material.html', {'form': form, 'error': "Falta configurar la API KEY en .env"})
                
                cantidad = form.cleaned_data['cantidad']
                porc_facil = form.cleaned_data['porc_facil']
                porc_intermedio = form.cleaned_data['porc_intermedio']
                porc_dificil = form.cleaned_data['porc_dificil']
                
                tipos = []
                if form.cleaned_data['incluir_om']: tipos.append("OM")
                if form.cleaned_data['incluir_vf']: tipos.append("VF")
                if form.cleaned_data['incluir_comp']: tipos.append("COMP")
                if form.cleaned_data['incluir_corr']: tipos.append("CORR")
                if not tipos: tipos = ["OM"]
                
                tipos_str = ", ".join(tipos)
                cant_facil = round(cantidad * (porc_facil / 100))
                cant_intermedio = round(cantidad * (porc_intermedio / 100))
                cant_dificil = cantidad - cant_facil - cant_intermedio

                client = OpenAI(api_key=api_key, base_url="https://api.groq.com/openai/v1")
                
                prompt_sistema = f"""
                Eres un profesor experto. A partir del texto, genera exactamente {cantidad} preguntas.
                Niveles de dificultad requeridos:
                - {cant_facil} de nivel 'FACIL'
                - {cant_intermedio} de nivel 'INTERMEDIO'
                - {cant_dificil} de nivel 'DIFICIL'
                Tipos de pregunta permitidos: {tipos_str}.
                REGLAS: OM (opciones, correcta 0-3), VF (correcta_bool), COMP (palabra), CORR (pares {{"A":"B"}}).
                Devuelve JSON: {{"preguntas": [{{"tipo": "OM", "nivel": "FACIL", "enunciado": "...", "opciones": [], "correcta": 0, "correcta_bool": null, "palabra": null, "pares": null}}]}}
                """
                
                respuesta = client.chat.completions.create(
                    model="openai/gpt-oss-20b", response_format={ "type": "json_object" }, 
                    messages=[{"role": "system", "content": prompt_sistema}, {"role": "user", "content": texto_recortado}]
                )
                
                datos_ia = json.loads(respuesta.choices[0].message.content)
                preguntas_creadas = []
                
                for item in datos_ia.get('preguntas', []):
                    tipo = item.get('tipo', 'OM')
                    nivel = item.get('nivel', 'FACIL')
                    enunciado = item.get('enunciado', 'Pregunta sin título')
                    contenido = {}
                    if tipo == 'VF': contenido = {"correcta": item.get('correcta_bool', True)}
                    elif tipo == 'COMP': contenido = {"palabra": item.get('palabra', '')}
                    elif tipo == 'CORR': contenido = {"pares": item.get('pares', {})}
                    
                    p = Pregunta.objects.create(
                        curso=curso_final, tipo=tipo, nivel=nivel, enunciado=enunciado, 
                        contenido=contenido, creado_por=request.user, anio=anio_seleccionado, seccion=seccion_seleccionada, es_ia=True 
                    )
                    
                    if tipo == 'OM':
                        for idx, opt_texto in enumerate(item.get('opciones', [])):
                            Opcion.objects.create(pregunta=p, texto=opt_texto, es_correcta=(idx == item.get('correcta', 0)), orden=idx)
                            
                    preguntas_creadas.append(p)
                    
                messages.success(request, f"¡Éxito! Se generaron {len(preguntas_creadas)} preguntas con IA.")
                
                preguntas_renderizadas = [{'obj': p, 'acierto': True} for p in preguntas_creadas]
                return render(request, 'core/examen_generado.html', {
                    'preguntas_resueltas': preguntas_renderizadas, 'nota_final': 20, 
                    'puntaje_obtenido': len(preguntas_creadas), 'total_preguntas': len(preguntas_creadas),
                    'curso_obj': curso_final, 'anio_obj': anio_seleccionado, 'titulo_examen': nombre_examen
                })
                
            except Exception as e:
                return render(request, 'core/subir_material.html', {'form': form, 'error': f"Error de la IA: {str(e)}"})
    else: form = SubirArchivoForm()
    return render(request, 'core/subir_material.html', {'form': form})

@login_required
@docente_required
def bancos_preguntas(request):
    carpeta_bancos = os.path.join(settings.MEDIA_ROOT, 'bancos_preguntas')
    if not os.path.exists(carpeta_bancos): os.makedirs(carpeta_bancos)
    archivos = [f for f in os.listdir(carpeta_bancos) if f.endswith(('.xlsx', '.xls'))]
    stats = {'facil': Pregunta.objects.filter(nivel='FACIL').count(), 'intermedio': Pregunta.objects.filter(nivel='INTERMEDIO').count(), 'dificil': Pregunta.objects.filter(nivel='DIFICIL').count()}
    examenes_creados = ExamenAsignado.objects.all().order_by('-fecha_creacion')
    context = {'archivos': archivos, 'stats': stats, 'cursos': Curso.objects.all().order_by('nombre'), 'anios': Anio.objects.all().order_by('numero'), 'secciones': Seccion.objects.all().order_by('letra'), 'examenes_creados': examenes_creados}
    return render(request, 'core/bancos_preguntas.html', context)

@login_required
@docente_required
def subir_banco_excel(request):
    if request.method == 'POST' and request.FILES.get('archivo_excel'):
        archivo = request.FILES['archivo_excel']
        carpeta_bancos = os.path.join(settings.MEDIA_ROOT, 'bancos_preguntas')
        ruta_archivo = os.path.join(carpeta_bancos, archivo.name)
        with open(ruta_archivo, 'wb+') as destination:
            for chunk in archivo.chunks(): destination.write(chunk)
        messages.success(request, f"Archivo {archivo.name} subido correctamente.")
    return redirect('bancos_preguntas')

@login_required
@docente_required
def procesar_banco_excel(request, nombre_archivo):
    if request.method != 'POST': return redirect('bancos_preguntas')
    ruta_archivo = os.path.join(settings.MEDIA_ROOT, 'bancos_preguntas', nombre_archivo)
    if not os.path.exists(ruta_archivo):
        messages.error(request, "El archivo no existe.")
        return redirect('bancos_preguntas')
    try:
        wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
        nombre_hoja_preguntas = next((s for s in wb.sheetnames if 'pregunta' in s.strip().lower()), wb.sheetnames[0])
        curso_id, anio_id, seccion_id = request.POST.get('curso'), request.POST.get('anio'), request.POST.get('seccion')
        if not curso_id or not anio_id or not seccion_id:
            messages.error(request, "Debes seleccionar Curso, Año y Sección.")
            return redirect('bancos_preguntas')
        curso, anio, seccion = Curso.objects.get(id=curso_id), Anio.objects.get(id=anio_id), Seccion.objects.get(id=seccion_id)
        hoja_preguntas = wb[nombre_hoja_preguntas]
        filas_procesadas = 0
        for row in hoja_preguntas.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 5 or not row[1] or not row[4]: continue
            tipo_excel, nivel_excel, enunciado = str(row[1]).strip().upper(), str(row[2]).strip().upper() if row[2] else 'FACIL', str(row[4]).strip()
            mapa_tipos = {'OPCION_MULTIPLE': 'OM', 'CORRELACION': 'CORR', 'VERDADERO_FALSO': 'VF', 'COMPLETAR': 'COMP', 'LIBRE': 'LIBRE'}
            tipo_django = mapa_tipos.get(tipo_excel, 'OM')
            if not Pregunta.objects.filter(enunciado=enunciado, curso=curso).exists():
                pregunta = Pregunta.objects.create(curso=curso, tipo=tipo_django, nivel=nivel_excel, enunciado=enunciado, contenido={}, creado_por=request.user, anio=anio, seccion=seccion, origen_banco=nombre_archivo)
                if tipo_django == 'OM':
                    correcta_letra = str(row[11]).strip().upper() if len(row) > 11 and row[11] else 'A'
                    for idx, txt in enumerate(row[6:11]):
                        if txt: Opcion.objects.create(pregunta=pregunta, texto=str(txt), es_correcta=(['A','B','C','D','E'][idx] == correcta_letra), orden=idx)
                elif tipo_django == 'VF': pregunta.contenido = {"correcta": str(row[11]).strip().upper() == 'V' if len(row) > 11 and row[11] else True}; pregunta.save()
                elif tipo_django == 'CORR':
                    pares = {}
                    for i in range(5):
                        idx_a, idx_b = 12+(i*2), 13+(i*2)
                        if len(row) > idx_b and row[idx_a] and row[idx_b]: pares[str(row[idx_a])] = str(row[idx_b])
                    pregunta.contenido = {"pares": pares}; pregunta.save()
                elif tipo_django == 'COMP': pregunta.contenido = {"palabra": str(row[11]).strip() if len(row) > 11 and row[11] else ''}; pregunta.save()
                elif tipo_django == 'LIBRE': pregunta.contenido = {"respuesta_modelo": str(row[11]).strip() if len(row) > 11 and row[11] else ''}; pregunta.save()
                filas_procesadas += 1
        messages.success(request, f"¡Éxito! Se cargaron {filas_procesadas} preguntas.")
    except Exception as e: messages.error(request, f"Error: {str(e)}")
    return redirect('bancos_preguntas')

@login_required
@docente_required
def eliminar_banco_excel(request, nombre_archivo):
    ruta_archivo = os.path.join(settings.MEDIA_ROOT, 'bancos_preguntas', nombre_archivo)
    if os.path.exists(ruta_archivo): os.remove(ruta_archivo)
    Pregunta.objects.filter(origen_banco=nombre_archivo).delete()
    messages.success(request, f"Archivo '{nombre_archivo}' eliminado.")
    return redirect('bancos_preguntas')

@login_required
@docente_required
def ver_preguntas_banco(request, nombre_archivo):
    preguntas = Pregunta.objects.filter(origen_banco=nombre_archivo).order_by('id')
    page_obj = Paginator(preguntas, 15).get_page(request.GET.get('page'))
    return render(request, 'core/ver_preguntas_banco.html', {'nombre_archivo': nombre_archivo, 'preguntas': page_obj})

@login_required
@docente_required
def generar_examen_porcentajes(request):
    if request.method == 'POST':
        curso_id, anio_id = request.POST.get('curso'), request.POST.get('anio')
        total_preguntas = int(request.POST.get('total_preguntas', 20))
        porc_facil, porc_intermedio = int(request.POST.get('porc_facil', 70)), int(request.POST.get('porc_intermedio', 20))
        banco_filtro, tiempo_minutos = request.POST.get('banco_filtro', ''), int(request.POST.get('tiempo_minutos', 30))
        guardar_examen = request.POST.get('guardar_examen') == 'on'
        cant_facil, cant_intermedio = round(total_preguntas * (porc_facil / 100)), round(total_preguntas * (porc_intermedio / 100))
        cant_dificil = total_preguntas - cant_facil - cant_intermedio
        query = Pregunta.objects.filter(curso_id=curso_id, anio_id=anio_id)
        if banco_filtro: query = query.filter(origen_banco=banco_filtro)
        def extraer(nivel, cant):
            disp = list(query.filter(nivel=nivel))
            return random.sample(disp, min(cant, len(disp))) if disp else []
        preguntas_examen = extraer('FACIL', cant_facil) + extraer('INTERMEDIO', cant_intermedio) + extraer('DIFICIL', cant_dificil)
        random.shuffle(preguntas_examen)
        if not preguntas_examen:
            messages.error(request, "No se encontraron preguntas.")
            return redirect('bancos_preguntas')
        if guardar_examen:
            curso, anio = Curso.objects.get(id=curso_id), Anio.objects.get(id=anio_id)
            nuevo_examen = ExamenAsignado.objects.create(titulo=f"Examen de {curso.nombre} - {anio}", curso=curso, anio=anio, tiempo_minutos=tiempo_minutos, creado_por=request.user)
            nuevo_examen.preguntas.set(preguntas_examen)
            messages.success(request, "¡Examen asignado exitosamente!")
            return redirect('bancos_preguntas')
        distractores = ["Ninguna de las anteriores", "Opción incorrecta", "Concepto no relacionado", "Definición errónea"]
        preguntas_renderizadas = [{'obj': p, 'opciones_corr': random.sample(list(p.contenido.get('pares', {}).values()) + [random.choice(distractores)], min(5, len(p.contenido.get('pares', {})) + 1)) if p.tipo.startswith('CORR') else []} for p in preguntas_examen]
        return render(request, 'core/rendir_examen.html', {'preguntas': preguntas_renderizadas, 'curso_obj': Curso.objects.get(id=curso_id), 'anio_obj': Anio.objects.get(id=anio_id), 'tiempo_minutos': tiempo_minutos})
    return redirect('bancos_preguntas')

@login_required
@docente_required
def ver_resultados_examen(request, examen_id):
    examen = get_object_or_404(ExamenAsignado, id=examen_id)
    alumnos = User.objects.filter(groups__name='Alumnos')
    resultados = ResultadoExamen.objects.filter(examen_asignado=examen)
    lista_alumnos = [{'alumno': a, 'resultado': resultados.filter(alumno=a).first()} for a in alumnos]
    return render(request, 'core/resultados_asignados.html', {'examen': examen, 'lista_alumnos': lista_alumnos})

@login_required
@docente_required
def ver_reportes(request):
    query = request.GET.get('q', '')
    preguntas_list = Pregunta.objects.filter(enunciado__icontains=query).order_by('-fecha_creacion') if query else Pregunta.objects.all().order_by('-fecha_creacion')
    preguntas = Paginator(preguntas_list, 10).get_page(request.GET.get('page'))
    tipos_data = {'labels': ['Opción Múltiple', 'Correlaciones', 'Verdadero/Falso', 'Completar'], 'values': [Pregunta.objects.filter(tipo='OM').count(), Pregunta.objects.filter(tipo__startswith='CORR').count(), Pregunta.objects.filter(tipo='VF').count(), Pregunta.objects.filter(tipo='COMP').count()]}
    notas_por_curso = ResultadoExamen.objects.values('curso__nombre').annotate(promedio=Avg('nota')).order_by('-promedio')[:5]
    notas_data = {'labels': [item['curso__nombre'] for item in notas_por_curso], 'values': [round(item['promedio'], 2) for item in notas_por_curso]}
    context = {'total': Pregunta.objects.count(), 'count_om': Pregunta.objects.filter(tipo='OM').count(), 'count_corr': Pregunta.objects.filter(tipo__startswith='CORR').count(), 'count_vf': Pregunta.objects.filter(tipo='VF').count(), 'count_comp': Pregunta.objects.filter(tipo='COMP').count(), 'count_libre': Pregunta.objects.filter(tipo='LIBRE').count(), 'especialidades': Especialidad.objects.all(), 'cursos': Curso.objects.all(), 'anios': Anio.objects.all(), 'preguntas': preguntas, 'query': query, 'tipos_data': tipos_data, 'notas_data': notas_data}
    return render(request, 'core/reportes.html', context)

@login_required
@docente_required
def eliminar_pregunta(request, pregunta_id):
    get_object_or_404(Pregunta, id=pregunta_id).delete()
    messages.success(request, "Pregunta eliminada exitosamente.")
    return redirect('ver_reportes')

@login_required
@docente_required
def editar_pregunta(request, pregunta_id):
    pregunta = get_object_or_404(Pregunta, id=pregunta_id)
    if request.method == 'POST':
        form = PreguntaForm(request.POST, instance=pregunta)
        if form.is_valid():
            p = form.save(commit=False); p.nivel = request.POST.get('nivel', 'FACIL'); p.save()
            messages.success(request, "Pregunta actualizada correctamente.")
            return redirect('ver_reportes')
    else: form = PreguntaForm(instance=pregunta)
    return render(request, 'core/editar_pregunta.html', {'form': form, 'pregunta': pregunta})

@login_required
@docente_required
def generar_examen(request):
    curso_id, anio_id = request.GET.get('curso'), request.GET.get('anio')
    preguntas_examen = list(Pregunta.objects.filter(curso_id=curso_id, anio_id=anio_id)[:20])
    if not preguntas_examen: return render(request, 'core/examen_generado.html', {'error': 'No hay preguntas.'})
    distractores = ["Ninguna de las anteriores", "Opción incorrecta", "Concepto no relacionado", "Definición errónea"]
    preguntas_renderizadas = [{'obj': p, 'opciones_corr': random.sample(list(p.contenido.get('pares', {}).values()) + [random.choice(distractores)], min(5, len(p.contenido.get('pares', {})) + 1)) if p.tipo.startswith('CORR') else []} for p in preguntas_examen]
    return render(request, 'core/examen_generado.html', {'preguntas_resueltas': [{'obj': p, 'acierto': True} for p in preguntas_examen], 'nota_final': 20, 'puntaje_obtenido': len(preguntas_examen), 'total_preguntas': len(preguntas_examen), 'curso_obj': Curso.objects.get(id=curso_id), 'anio_obj': Anio.objects.get(id=anio_id)})

@login_required
@docente_required
def generar_examen_manual(request):
    if request.method == 'POST':
        pregunta_ids = request.POST.getlist('preguntas_seleccionadas')
        if not pregunta_ids:
            messages.error(request, "No seleccionaste ninguna pregunta.")
            return redirect('ver_reportes')
        preguntas_examen = list(Pregunta.objects.filter(id__in=pregunta_ids))
        return render(request, 'core/examen_generado.html', {'preguntas_resueltas': [{'obj': p, 'acierto': True} for p in preguntas_examen], 'nota_final': 20, 'puntaje_obtenido': len(preguntas_examen), 'total_preguntas': len(preguntas_examen), 'curso_obj': None, 'anio_obj': None})
    return redirect('ver_reportes')

@login_required
@docente_required
def ver_historial_notas(request):
    return render(request, 'core/historial_notas.html', {'notas': ResultadoExamen.objects.all().order_by('-fecha_rendido')})

@login_required
def rendir_examen_asignado(request, examen_id):
    examen = get_object_or_404(ExamenAsignado, id=examen_id)
    preguntas_examen = list(examen.preguntas.all())
    random.shuffle(preguntas_examen)
    distractores = ["Ninguna de las anteriores", "Opción incorrecta", "Concepto no relacionado", "Definición errónea"]
    preguntas_renderizadas = [{'obj': p, 'opciones_corr': random.sample(list(p.contenido.get('pares', {}).values()) + [random.choice(distractores)], min(5, len(p.contenido.get('pares', {})) + 1)) if p.tipo.startswith('CORR') else []} for p in preguntas_examen]
    return render(request, 'core/rendir_examen.html', {'preguntas': preguntas_renderizadas, 'curso_obj': examen.curso, 'anio_obj': examen.anio, 'tiempo_minutos': examen.tiempo_minutos, 'examen_asignado_id': examen.id})

@login_required
def rendir_examen(request):
    curso_id, anio_id = request.GET.get('curso'), request.GET.get('anio')
    preguntas_examen = list(Pregunta.objects.filter(curso_id=curso_id, anio_id=anio_id)[:20])
    if not preguntas_examen: return render(request, 'core/examen_generado.html', {'error': 'No hay preguntas.'})
    distractores = ["Ninguna de las anteriores", "Opción incorrecta", "Concepto no relacionado", "Definición errónea"]
    preguntas_renderizadas = [{'obj': p, 'opciones_corr': random.sample(list(p.contenido.get('pares', {}).values()) + [random.choice(distractores)], min(5, len(p.contenido.get('pares', {})) + 1)) if p.tipo.startswith('CORR') else []} for p in preguntas_examen]
    return render(request, 'core/rendir_examen.html', {'preguntas': preguntas_renderizadas, 'curso_obj': Curso.objects.filter(id=curso_id).first(), 'anio_obj': Anio.objects.filter(id=anio_id).first(), 'tiempo_minutos': 30})

@login_required
def calificar_examen(request):
    if request.method == 'POST':
        pregunta_ids = request.POST.getlist('pregunta_ids')
        examen_asignado_id = request.POST.get('examen_asignado_id') 
        puntaje, resultados, preguntas_libres_guardadas = 0, [], []
        for p_id in pregunta_ids:
            p = Pregunta.objects.get(id=p_id)
            correcta, respuesta_alumno, respuesta_correcta = False, "", ""
            if p.tipo == 'OM':
                opcion_id = request.POST.get(f'om_{p.id}')
                if opcion_id:
                    opcion_sel = Opcion.objects.get(id=opcion_id)
                    respuesta_alumno = opcion_sel.texto
                    opcion_correcta = p.opciones.filter(es_correcta=True).first()
                    if opcion_correcta:
                        respuesta_correcta = opcion_correcta.texto
                        if opcion_sel.es_correcta: correcta = True
            elif p.tipo.startswith('CORR'):
                aciertos_corr, pares = 0, p.contenido.get('pares', {})
                for i, (clave, valor_correcto) in enumerate(pares.items()):
                    resp = request.POST.get(f'corr_{p.id}_{i}', '').strip().lower()
                    if resp == valor_correcto.strip().lower(): aciertos_corr += 1
                if aciertos_corr == len(pares) and len(pares) > 0: correcta = True
                respuesta_alumno, respuesta_correcta = f"{aciertos_corr}/{len(pares)} correctas", "Emparejamiento completo"
            elif p.tipo == 'VF':
                respuesta_alumno = request.POST.get(f'vf_{p.id}')
                es_vf_correcta = p.contenido.get('correcta', False)
                respuesta_correcta = "Verdadero" if es_vf_correcta else "Falso"
                if (respuesta_alumno == 'true' and es_vf_correcta) or (respuesta_alumno == 'false' and not es_vf_correcta): correcta = True
            elif p.tipo == 'COMP':
                respuesta_alumno = request.POST.get(f'comp_{p.id}', '').strip().lower()
                respuesta_correcta = p.contenido.get('palabra', '')
                if respuesta_alumno == respuesta_correcta.strip().lower(): correcta = True
            elif p.tipo == 'LIBRE':
                respuesta_alumno, respuesta_correcta, correcta = request.POST.get(f'libre_{p.id}', ''), "Revisión Manual", True
                preguntas_libres_guardadas.append({'pregunta_id': p.id, 'enunciado': p.enunciado, 'respuesta_alumno': respuesta_alumno, 'puntaje_asignado': 1.0, 'revisado': False})
            if correcta: puntaje += 1
            resultados.append({'pregunta': p, 'respuesta_alumno': respuesta_alumno, 'respuesta_correcta': respuesta_correcta, 'acierto': correcta})
        nota_final = (puntaje / len(pregunta_ids)) * 20 if pregunta_ids else 0
        if pregunta_ids:
            primera = Pregunta.objects.get(id=pregunta_ids[0])
            ResultadoExamen.objects.create(alumno=request.user, curso=primera.curso, anio=primera.anio, seccion=primera.seccion, nota=round(nota_final, 2), examen_asignado_id=examen_asignado_id if examen_asignado_id else None, respuestas_detalle={'libres': preguntas_libres_guardadas, 'total_preguntas': len(pregunta_ids)})
        return render(request, 'core/resultado_examen.html', {'resultados': resultados, 'nota_final': round(nota_final), 'puntaje': puntaje, 'total': len(pregunta_ids)})
    return redirect('home')

@login_required
@docente_required
def revisar_libres(request, resultado_id):
    resultado = get_object_or_404(ResultadoExamen, id=resultado_id)
    libres = resultado.respuestas_detalle.get('libres', [])
    return render(request, 'core/revisar_libres.html', {'resultado': resultado, 'libres': libres})

@login_required
@docente_required
def guardar_revision_libres(request, resultado_id):
    if request.method == 'POST':
        resultado = get_object_or_404(ResultadoExamen, id=resultado_id)
        libres = resultado.respuestas_detalle.get('libres', [])
        total_preguntas = resultado.respuestas_detalle.get('total_preguntas', 20)
        puntaje_anterior_libres, puntaje_nuevo_libres = 0, 0
        for i, libre in enumerate(libres):
            puntaje_anterior_libres += 1 
            nuevo_puntaje = float(request.POST.get(f'puntaje_{i}', 1))
            puntaje_nuevo_libres += nuevo_puntaje
            libre['puntaje_asignado'], libre['revisado'] = nuevo_puntaje, True
        puntaje_actual_obtenido = (resultado.nota * total_preguntas) / 20
        puntaje_nuevo_obtenido = puntaje_actual_obtenido - puntaje_anterior_libres + puntaje_nuevo_libres
        resultado.nota = round((puntaje_nuevo_obtenido / total_preguntas) * 20, 2)
        resultado.respuestas_detalle['libres'] = libres
        resultado.save()
        messages.success(request, "Calificación de preguntas libres actualizada correctamente.")
        return redirect('ver_historial_notas')
    return redirect('ver_historial_notas')

def estilizar_excel(ws, headers):
    header_font = Font(color="FFFFFF", bold=True, size=12)
    header_fill = PatternFill(start_color="0B2545", end_color="0B2545", fill_type="solid")
    alignment, border = Alignment(horizontal="center", vertical="center"), Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    ws.append(headers)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, alignment, border

@login_required
@docente_required
def exportar_preguntas_excel(request):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Banco de Preguntas"
    headers = ['ID', 'Tipo', 'Nivel', 'Especialidad', 'Curso', 'Año', 'Sección', 'Enunciado', 'Creado por', 'Fecha']
    estilizar_excel(ws, headers)
    for p in Pregunta.objects.all().order_by('id'):
        ws.append([p.id, p.get_tipo_display(), p.get_nivel_display(), p.curso.especialidad.nombre, p.curso.nombre, p.anio.get_numero_display(), p.seccion.letra, p.enunciado, p.creado_por.username if p.creado_por else "N/A", p.fecha_creacion.strftime("%d/%m/%Y %H:%M")])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(cell.value)) for cell in col if cell.value) + 2
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Banco_de_Preguntas_CITEN.xlsx"'
    wb.save(response); return response

@login_required
@docente_required
def exportar_notas_excel(request):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Historial de Notas"
    headers = ['Fecha', 'Alumno', 'Curso', 'Especialidad', 'Año', 'Sección', 'Nota Final']
    estilizar_excel(ws, headers)
    for n in ResultadoExamen.objects.all().order_by('-fecha_rendido'):
        ws.append([n.fecha_rendido.strftime("%d/%m/%Y %H:%M"), n.alumno.username, n.curso.nombre, n.curso.especialidad.nombre, n.anio.get_numero_display(), n.seccion.letra if n.seccion else "N/A", n.nota])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(cell.value)) for cell in col if cell.value) + 2
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Historial_de_Notas_CITEN.xlsx"'
    wb.save(response); return response